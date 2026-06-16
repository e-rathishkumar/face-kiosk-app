import csv

from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from sqlalchemy.orm import Session

from app.models.report_history import (
    ReportHistory
)

from app.repositories.report_repository import (
    ReportRepository
)

from app.repositories.attendance_repository import (
    AttendanceRepository
)

from app.repositories.recognition_log_repository import (
    RecognitionLogRepository
)

from app.repositories.unrecognized_repository import (
    UnrecognizedRepository
)

from app.services.audit_service import (
    AuditService
)


EXPORT_DIR = Path("exports")
EXPORT_DIR.mkdir(exist_ok=True)


class ExportService:

    @staticmethod
    def _save_history(
        db: Session,
        report_name: str,
        file_path
    ):
        history = ReportHistory(
            report_name=report_name,
            generated_by="system",
            file_url=str(file_path),
            generated_at=datetime.utcnow()
        )

        ReportRepository.create(
            db,
            history
        )

        AuditService.log(
            db=db,
            action="EXPORT",
            entity_type="REPORT",
            entity_id=report_name,
            new_value=str(file_path)
        )

    @staticmethod
    def export_attendance_csv(
        db: Session
    ):
        sessions = AttendanceRepository.get_all(db)

        file_path = (
            EXPORT_DIR /
            "attendance_report.csv"
        )

        with open(
            file_path,
            "w",
            newline=""
        ) as file:

            writer = csv.writer(file)

            writer.writerow([
                "ID",
                "Employee ID",
                "Kiosk ID",
                "Check In",
                "Check Out",
                "Status"
            ])

            for row in sessions:
                writer.writerow([
                    row.id,
                    row.employee_id,
                    row.kiosk_id,
                    row.check_in_time,
                    row.check_out_time,
                    row.status
                ])

        ExportService._save_history(
            db,
            "attendance_csv",
            file_path
        )

        return file_path

    @staticmethod
    def export_attendance_xlsx(
        db: Session
    ):
        sessions = AttendanceRepository.get_all(db)

        workbook = Workbook()
        sheet = workbook.active

        sheet.append([
            "ID",
            "Employee ID",
            "Kiosk ID",
            "Check In",
            "Check Out",
            "Status"
        ])

        for row in sessions:
            sheet.append([
                str(row.id),
                str(row.employee_id),
                str(row.kiosk_id),
                str(row.check_in_time),
                str(row.check_out_time),
                str(row.status)
            ])

        file_path = (
            EXPORT_DIR /
            "attendance_report.xlsx"
        )

        workbook.save(file_path)

        ExportService._save_history(
            db,
            "attendance_xlsx",
            file_path
        )

        return file_path

    @staticmethod
    def export_recognition_csv(
        db: Session
    ):
        logs = RecognitionLogRepository.get_all(db)

        file_path = (
            EXPORT_DIR /
            "recognition_report.csv"
        )

        with open(
            file_path,
            "w",
            newline=""
        ) as file:

            writer = csv.writer(file)

            writer.writerow([
                "ID",
                "Employee ID",
                "Kiosk ID",
                "Confidence",
                "Event Time"
            ])

            for row in logs:
                writer.writerow([
                    row.id,
                    row.employee_id,
                    row.kiosk_id,
                    row.confidence_score,
                    row.event_time
                ])

        ExportService._save_history(
            db,
            "recognition_csv",
            file_path
        )

        return file_path

    @staticmethod
    def export_recognition_xlsx(
        db: Session
    ):
        logs = RecognitionLogRepository.get_all(db)

        workbook = Workbook()
        sheet = workbook.active

        sheet.append([
            "ID",
            "Employee ID",
            "Kiosk ID",
            "Confidence",
            "Event Time"
        ])

        for row in logs:
            sheet.append([
                str(row.id),
                str(row.employee_id),
                str(row.kiosk_id),
                row.confidence_score,
                str(row.event_time)
            ])

        file_path = (
            EXPORT_DIR /
            "recognition_report.xlsx"
        )

        workbook.save(file_path)

        ExportService._save_history(
            db,
            "recognition_xlsx",
            file_path
        )

        return file_path

    @staticmethod
    def export_unrecognized_csv(
        db: Session
    ):
        entries = UnrecognizedRepository.get_all(db)

        file_path = (
            EXPORT_DIR /
            "unrecognized_report.csv"
        )

        with open(
            file_path,
            "w",
            newline=""
        ) as file:

            writer = csv.writer(file)

            writer.writerow([
                "ID",
                "Kiosk ID",
                "Confidence",
                "Status"
            ])

            for row in entries:
                writer.writerow([
                    row.id,
                    row.kiosk_id,
                    row.confidence_score,
                    row.status
                ])

        ExportService._save_history(
            db,
            "unrecognized_csv",
            file_path
        )

        return file_path

    @staticmethod
    def export_unrecognized_xlsx(
        db: Session
    ):
        entries = UnrecognizedRepository.get_all(db)

        workbook = Workbook()
        sheet = workbook.active

        sheet.append([
            "ID",
            "Kiosk ID",
            "Confidence",
            "Status"
        ])

        for row in entries:
            sheet.append([
                str(row.id),
                str(row.kiosk_id),
                row.confidence_score,
                str(row.status)
            ])

        file_path = (
            EXPORT_DIR /
            "unrecognized_report.xlsx"
        )

        workbook.save(file_path)

        ExportService._save_history(
            db,
            "unrecognized_xlsx",
            file_path
        )

        return file_path
